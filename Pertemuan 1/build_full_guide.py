import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()
    
    # -------------------------------------------------------------
    # Page Setup (A4 with 1-inch margins)
    # -------------------------------------------------------------
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4 width
        section.page_height = Inches(11.69) # A4 height
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header & Footer setup
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Modul PAM Pertemuan 1 • Panduan Instalasi Android Studio & Spesifikasi Laptop")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(140, 145, 150)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        frun = fp.add_run("Pengembangan Aplikasi Mobile (PAM) — Android Studio Quail 4 (2026.1.4)")
        frun.font.name = "Calibri"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = RGBColor(140, 145, 150)

    # Styling helper functions
    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        tcPr.append(shd)

    def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(
            f'<w:tcMar {nsdecls("w")}>'
            f'<w:top w:w="{top}" w:type="dxa"/>'
            f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
            f'<w:left w:w="{left}" w:type="dxa"/>'
            f'<w:right w:w="{right}" w:type="dxa"/>'
            f'</w:tcMar>'
        )
        tcPr.append(tcMar)

    def set_table_borders(table, color="D1D5DB", sz="4", val="single"):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:left w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
            f'<w:insideV w:val="none"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_callout(text, title="CATATAN PENTING", box_type="tip"):
        colors = {
            "tip": {"bg": "E8F5E9", "border": "2E7D32", "title": "2E7D32", "icon": "💡"},
            "note": {"bg": "E1F5FE", "border": "0277BD", "title": "0277BD", "icon": "ℹ️"},
            "warning": {"bg": "FFF8E1", "border": "F57F17", "title": "E65100", "icon": "⚠️"},
            "info": {"bg": "F3E5F5", "border": "7B1FA2", "title": "6A1B9A", "icon": "📌"}
        }
        cfg = colors.get(box_type, colors["note"])
        
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.27)
        
        set_cell_background(cell, cfg["bg"])
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:left w:val="single" w:sz="36" w:space="0" w:color="{cfg["border"]}"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)
        
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        run_t = p.add_run(f"{cfg['icon']} {title}")
        run_t.bold = True
        run_t.font.name = "Arial"
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = RGBColor.from_string(cfg["title"])
        
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(0)
        p2.paragraph_format.line_spacing = 1.15
        run_body = p2.add_run(text)
        run_body.font.name = "Calibri"
        run_body.font.size = Pt(10)
        run_body.font.color.rgb = RGBColor(40, 40, 40)
        
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)

    def add_figure(img_path, caption, width=Inches(5.4)):
        if not os.path.exists(img_path):
            print(f"Warning: Image {img_path} not found!")
            return
            
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(12)
        
        run_cap = p_cap.add_run(caption)
        run_cap.font.name = "Calibri"
        run_cap.font.size = Pt(9.5)
        run_cap.font.italic = True
        run_cap.font.color.rgb = RGBColor(90, 95, 100)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(7, 48, 66) # Android Navy
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(27, 85, 114) # Slate Blue/Teal
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(44, 62, 80)
        return p

    def add_p(text, bold_prefix="", space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.name = "Calibri"
            run_b.font.size = Pt(11)
            run_b.font.color.rgb = RGBColor(30, 30, 30)
        run_t = p.add_run(text)
        run_t.font.name = "Calibri"
        run_t.font.size = Pt(11)
        run_t.font.color.rgb = RGBColor(45, 50, 55)
        return p

    def add_bullet(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            run_b = p.add_run(bold_prefix)
            run_b.bold = True
            run_b.font.name = "Calibri"
            run_b.font.size = Pt(10.5)
            run_b.font.color.rgb = RGBColor(30, 30, 30)
        run_t = p.add_run(text)
        run_t.font.name = "Calibri"
        run_t.font.size = Pt(10.5)
        run_t.font.color.rgb = RGBColor(45, 50, 55)
        return p

    print("Building Document Structure...")

    # =========================================================================
    # HALAMAN COVER (SAMPUL DEPAN ELEGAN & AKADEMIK)
    # =========================================================================
    p_cov_top = doc.add_paragraph()
    p_cov_top.paragraph_format.space_before = Pt(10)
    p_cov_top.paragraph_format.space_after = Pt(4)
    run_cat = p_cov_top.add_run("MODUL PRAKTIKUM & PANDUAN TEKNIS LENGKAP")
    run_cat.bold = True
    run_cat.font.name = "Arial"
    run_cat.font.size = Pt(11)
    run_cat.font.color.rgb = RGBColor(46, 125, 50) # Android Accent Green

    p_cov_course = doc.add_paragraph()
    p_cov_course.paragraph_format.space_before = Pt(0)
    p_cov_course.paragraph_format.space_after = Pt(20)
    run_crs = p_cov_course.add_run("PENGEMBANGAN APLIKASI MOBILE (PAM) — PERTEMUAN 1")
    run_crs.bold = True
    run_crs.font.name = "Arial"
    run_crs.font.size = Pt(13)
    run_crs.font.color.rgb = RGBColor(7, 48, 66)

    # Decorative colored bar
    p_bar = doc.add_paragraph()
    p_bar.paragraph_format.space_before = Pt(0)
    p_bar.paragraph_format.space_after = Pt(24)
    run_bar = p_bar.add_run("━" * 48)
    run_bar.font.name = "Arial"
    run_bar.font.size = Pt(14)
    run_bar.font.color.rgb = RGBColor(61, 220, 132) # Android Robot Green

    # Main Title
    p_cov_title = doc.add_paragraph()
    p_cov_title.paragraph_format.space_before = Pt(0)
    p_cov_title.paragraph_format.space_after = Pt(12)
    p_cov_title.paragraph_format.line_spacing = 1.15
    run_title = p_cov_title.add_run("PANDUAN LENGKAP PENGINSTALAN ANDROID STUDIO BESERTA ANALISIS SPESIFIKASI SISTEM LAPTOP")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(7, 48, 66)

    # Subtitle
    p_cov_sub = doc.add_paragraph()
    p_cov_sub.paragraph_format.space_before = Pt(0)
    p_cov_sub.paragraph_format.space_after = Pt(36)
    p_cov_sub.paragraph_format.line_spacing = 1.2
    run_sub = p_cov_sub.add_run(
        "Panduan Komprehensif Berbasis Dokumentasi Visual (15 Langkah Lengkap): "
        "Mulai dari Persiapan Hardware, Pengunduhan Installer Resmi, Konfigurasi Inti Software, "
        "Setup Android SDK Toolchain, Lisensi, AVD Emulator, hingga Pengujian Proyek Perdana dan Troubleshooting."
    )
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(90, 95, 100)

    # Metadata Box (Card Table)
    meta_data = [
        ("Nama Mahasiswa", "Muhammad Fatahillah Farid"),
        ("Nomor Induk Mahasiswa (NIM)", "123140203"),
        ("Program Studi / Institusi", "Teknik Informatika — Institut Teknologi Sumatera (ITERA)"),
        ("Mata Kuliah & Pertemuan", "Pengembangan Aplikasi Mobile (PAM) — Pertemuan 1 (P1)"),
        ("Versi Android Studio", "Android Studio Quail 4 (Rilis Build 2026.1.4)"),
        ("Sistem Operasi Target", "Microsoft Windows 11 / Windows 10 (64-bit)"),
        ("Cakupan Laporan", "Analisis Hardware, 15 Langkah Setup Visual, Optimasi RAM 8 GB, & Troubleshooting")
    ]
    tbl_meta = doc.add_table(rows=len(meta_data), cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_meta, color="BDC3C7", sz="4")
    
    for row_idx, (k, v) in enumerate(meta_data):
        cell_k, cell_v = tbl_meta.rows[row_idx].cells
        cell_k.width = Inches(2.2)
        cell_v.width = Inches(4.07)
        
        set_cell_background(cell_k, "F4F6F7")
        set_cell_background(cell_v, "FFFFFF")
        set_cell_margins(cell_k, top=100, bottom=100, left=140, right=140)
        set_cell_margins(cell_v, top=100, bottom=100, left=140, right=140)
        
        pk = cell_k.paragraphs[0]
        pk.paragraph_format.space_before = Pt(0)
        pk.paragraph_format.space_after = Pt(0)
        rk = pk.add_run(k)
        rk.bold = True
        rk.font.name = "Arial"
        rk.font.size = Pt(9.5)
        rk.font.color.rgb = RGBColor(44, 62, 80)
        
        pv = cell_v.paragraphs[0]
        pv.paragraph_format.space_before = Pt(0)
        pv.paragraph_format.space_after = Pt(0)
        rv = pv.add_run(v)
        rv.font.name = "Calibri"
        rv.font.size = Pt(9.5)
        rv.font.color.rgb = RGBColor(50, 50, 50)

    doc.add_page_break()

    # =========================================================================
    # KATA PENGANTAR & TUJUAN PEMBELAJARAN
    # =========================================================================
    add_h1("KATA PENGANTAR & TUJUAN PEMBELAJARAN")
    add_p(
        "Puji syukur kita panjatkan ke hadirat Tuhan Yang Maha Esa atas tersusunnya dokumen panduan praktikum "
        "Pengembangan Aplikasi Mobile (PAM) Pertemuan 1. Modul ini disusun secara khusus untuk memberikan asistensi "
        "teknis yang menyeluruh bagi mahasiswa semester 7 dalam menyiapkan lingkungan pengembangan (Integrated Development "
        "Environment / IDE) Android Studio versi Quail 4 (2026.1.4)."
    )
    add_p(
        "Lingkungan pengembangan aplikasi mobile merupakan pondasi paling fundamental dalam seluruh rangkaian perkuliahan PAM. "
        "Ketidaksesuaian spesifikasi hardware atau kesalahan langkah pada konfigurasi awal sering kali menimbulkan hambatan fatal, "
        "seperti proses kompilasi Gradle yang terhenti (stuck), emulator gagal memuat (crash), hingga crash kehabisan memori (Out of Memory). "
        "Oleh karena itu, modul ini tidak hanya memaparkan tata cara penginstalan langkah-demi-langkah dengan dokumentasi visual otentik, "
        "tetapi juga mengupas tuntas standarisasi kebutuhan perangkat keras, analisis alokasi sumber daya sistem, strategi optimasi laptop "
        "berspesifikasi pas-pasan, serta panduan troubleshooting teruji."
    )
    
    add_h2("Tujuan Instruksional Praktikum")
    add_bullet("Memahami arsitektur internal Android Studio, Android SDK (Software Development Kit), dan Gradle Build System.", "1. Pemahaman Konseptual: ")
    add_bullet("Mampu mengidentifikasi kesesuaian spesifikasi laptop pribadi dengan kebutuhan beban kerja pengembangan aplikasi mobile Android modern.", "2. Analisis Kebutuhan Sistem: ")
    add_bullet("Mampu melaksanakan pengunduhan, penginstalan, dan konfigurasi Android Studio Quail 4 secara mandiri dan sistematis tanpa ada komponen yang terlewat.", "3. Keterampilan Instalasi: ")
    add_bullet("Mampu mengatur lisensi SDK, komponen Build-Tools, Platform-Tools, dan emulator Android Virtual Device (AVD).", "4. Konfigurasi Toolchain: ")
    add_bullet("Mampu mendiagnosis dan mengatasi permasalahan teknis umum (seperti aktivasi VT-x/AMD-V BIOS, error dependensi Gradle, dan optimasi alokasi RAM).", "5. Kemandirian Troubleshooting: ")

    add_callout(
        "Seluruh tahapan dalam dokumen ini telah divalidasi dan diuji secara langsung pada lingkungan sistem operasi Windows 64-bit "
        "dengan memanfaatkan berkas installer Android Studio Quail 4 (2026.1.4). Ikuti urutan instruksi secara berurutan dan hindari "
        "melewatkan langkah apa pun demi menjamin kelancaran praktikum Anda.",
        title="PETUNJUK BAGI MAHASISWA",
        box_type="tip"
    )

    # =========================================================================
    # BAB I: PENGENALAN ANDROID STUDIO
    # =========================================================================
    add_h1("BAB I: PENGENALAN ANDROID STUDIO")
    
    add_h2("1.1 Apa Itu Android Studio?")
    add_p(
        "Android Studio adalah Lingkungan Pengembangan Terpadu (Integrated Development Environment atau IDE) resmi "
        "yang dirancang dan didistribusikan secara gratis oleh Google untuk pengembangan sistem operasi Android. "
        "Diperkenalkan pertama kali pada Google I/O tahun 2013 sebagai pengganti Eclipse Android Development Tools (ADT), "
        "Android Studio dibangun di atas basis perangkat lunak IntelliJ IDEA buatan JetBrains yang telah teruji kehandalannya "
        "dalam menangani proyek perangkat lunak skala enterprise."
    )
    add_p(
        "Dalam perkembangan terkini, versi yang digunakan pada modul ini adalah Android Studio rilis kanonis Quail 4 (versi 2026.1.4). "
        "Versi ini membawa lompatan teknologi yang signifikan, terutama dalam integrasi kecerdasan buatan Gemini AI Assistant, "
        "dukungan penuh Android SDK Platform 37, peningkatan drastis kecepatan kompilasi Gradle dengan Kotlin Multiplatform, "
        "serta perbaikan arsitektur Android Virtual Device (AVD) yang jauh lebih hemat daya dan responsif."
    )

    add_h2("1.2 Arsitektur Komponen Utama Android Studio")
    add_p(
        "Agar tidak terjadi kerancuan antara aplikasi editor dan perkakas pendukungnya, mahasiswa wajib memahami lima pilar "
        "komponen utama yang membentuk ekosistem Android Studio:"
    )
    add_bullet(
        "Perangkat lunak antarmuka grafis yang menyediakan editor teks pintar, inspeksi kode statis (Lint), refactoring otomatis, "
        "analisis penggunaan memori (Memory Profiler), Network Inspector, dan integrasi kontrol versi Git.",
        "a. Android Studio IDE: "
    )
    add_bullet(
        "Paket pustaka dan utilitas kompilasi yang memungkinkan kode sumber Java/Kotlin diterjemahkan menjadi aplikasi Android. Terdiri dari:\n"
        "   • SDK Platform: Berisi pustaka API bawaan Android (misal Android 37.0) yang menjadi acuan framework aplikasi.\n"
        "   • SDK Build-Tools (versi 36.0.0): Kumpulan utilitas biner (D8 compiler, AAPT2 packaging, R8 code shrinker) yang bertugas merangkai kode dan aset menjadi format .APK atau .AAB.\n"
        "   • Platform-Tools (ADB - Android Debug Bridge): Antarmuka baris perintah yang menghubungkan PC pengembang dengan perangkat Android fisik atau virtual.",
        "b. Android SDK (Software Development Kit): "
    )
    add_bullet(
        "Mesin otomasi pembangun proyek (build automation tool) tingkat lanjut. Gradle mengatur manajemen pustaka eksternal "
        "(libraries dependency), konfigurasi varian build (Debug vs Release), serta kompilasi modular lintas perangkat.",
        "c. Gradle Build System: "
    )
    add_bullet(
        "Simulator perangkat ponsel cerdas virtual berbasis emulator QEMU dengan akselerasi perangkat keras (Hardware Virtualization). "
        "Emulator ini memungkinkan pengembang menjalankan dan menguji aplikasi secara real-time di layar komputer tanpa membutuhkan HP fisik.",
        "d. Android Virtual Device (AVD Emulator): "
    )
    add_bullet(
        "Kecerdasan buatan generatif Google yang terintegrasi secara langsung di dalam IDE untuk menjawab pertanyaan teknis, "
        "menghasilkan potongan kode secara otomatis, mendeteksi potensi memory leak, dan memberikan saran perbaikan error kompilasi secara instan.",
        "e. Gemini in Android Studio: "
    )

    # =========================================================================
    # BAB II: SPESIFIKASI KEBUTUHAN SISTEM LAPTOP / PC
    # =========================================================================
    add_h1("BAB II: SPESIFIKASI KEBUTUHAN SISTEM LAPTOP / PC")

    add_h2("2.1 Mengapa Android Studio Memerlukan Spesifikasi Hardware Tinggi?")
    add_p(
        "Banyak mahasiswa pemula yang mengeluhkan bahwa laptop mereka mendadak menjadi sangat panas, lambat, atau bahkan "
        "mengalami 'freeze / Not Responding' saat membuka Android Studio. Hal ini terjadi karena Android Studio bukanlah "
        "sekadar text editor ringan seperti Notepad atau Visual Studio Code biasa. Saat Anda bekerja di Android Studio, komputer "
        "menjalankan setidaknya empat mesin komputasi besar secara serentak:"
    )
    add_bullet("Menjalankan Java Virtual Machine (JVM) internal untuk melayani indexing ribuan berkas class, auto-completion, dan analisis error realtime.", "1. Proses IDE IntelliJ: ")
    add_bullet("Berjalan sebagai background daemon mandiri di luar IDE yang mengonsumsi memori besar untuk memanipulasi graph dependency dan mengompilasi kode.", "2. Gradle Daemon: ")
    add_bullet("Menjalankan sistem operasi Android lengkap (Linux kernel + Android runtime) di atas hypervisor virtualisasi laptop.", "3. Android Virtual Device (AVD): ")
    add_bullet("Membuka puluhan tab browser untuk mencari dokumentasi API, StackOverflow, atau materi perkuliahan.", "4. Web Browser & Tool Eksternal: ")

    add_h2("2.2 Tabel Komparasi Spesifikasi Laptop (System Requirements)")
    add_p(
        "Berikut adalah tabel perbandingan komprehensif antara spesifikasi minimum resmi dari Google, spesifikasi yang direkomendasikan "
        "untuk praktikum standar, dan spesifikasi optimal untuk pengalaman pengembangan aplikasi tanpa hambatan:"
    )

    # Comparison Table
    tbl_spec = doc.add_table(rows=9, cols=4)
    tbl_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_spec, color="BDC3C7", sz="4")
    
    headers = ["Komponen Hardware", "Minimum Resmi (Google)", "Rekomendasi Standar", "Rekomendasi Optimal (PAM)"]
    for i, h in enumerate(headers):
        cell = tbl_spec.rows[0].cells[i]
        set_cell_background(cell, "073042") # Deep Navy
        set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    specs_data = [
        ("Sistem Operasi", "Windows 8 / 10 / 11 (64-bit)", "Windows 10 / 11 (64-bit) Update Terkini", "Windows 11 (64-bit) Pro / Home"),
        ("Processor (CPU)", "x86_64 CPU; Intel Core Gen 2 / AMD dengan Hypervisor", "Intel Core i5 (Gen 10+) / AMD Ryzen 5 (Seri 4000/5000) (4 Core/8 Thread)", "Intel Core i7 (Gen 12+) / AMD Ryzen 7 (Seri 6000+) (8-16 Core/Thread)"),
        ("Virtualisasi", "VT-x / AMD-V wajib ada di CPU", "VT-x / SVM diaktifkan di BIOS & WHPX aktif", "Hardware Virtualization + Hyper-V aktif penuh"),
        ("RAM (Memori)", "8 GB RAM fisik", "16 GB DDR4 / DDR5 Dual Channel", "32 GB DDR4 / DDR5 Dual Channel"),
        ("Storage (Disk)", "8 GB ruang kosong (HDD diizinkan)", "SSD SATA / NVMe minimal 30 GB ruang kosong", "SSD NVMe M.2 PCIe Gen 4 minimal 60-100 GB kosong"),
        ("Resolusi Layar", "1280 x 800 piksel", "1920 x 1080 piksel (Full HD 1080p)", "1920 x 1080 atau 2K/4K Dual Monitor"),
        ("Kartu Grafis (GPU)", "Integrated GPU (OpenGL 3.3)", "Dedicated GPU (NVIDIA GTX/RTX atau AMD Radeon)", "Dedicated GPU RTX 3050+ (DirectX 12 / Vulkan support)"),
        ("Koneksi Internet", "Koneksi 2-5 Mbps (Stabil)", "Koneksi 10-20 Mbps (Kuota min. 5 GB)", "Koneksi 50+ Mbps Tanpa Kuota (Fiber Optic)")
    ]

    for row_idx, row_data in enumerate(specs_data, start=1):
        bg = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = tbl_spec.rows[row_idx].cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=110, right=110)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(8.5)
            if col_idx == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(7, 48, 66)
            elif col_idx == 3:
                run.font.color.rgb = RGBColor(46, 125, 50)
            else:
                run.font.color.rgb = RGBColor(50, 50, 50)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_after = Pt(8)

    add_h2("2.3 Pembahasan Rinci Komponen Kritis")
    
    add_h3("A. Kapasitas RAM: Mengapa 8 GB Sangat Mepet?")
    add_p(
        "Meskipun Google secara resmi menuliskan batas minimum 8 GB RAM, dalam realitas praktikum laboratorium, 8 GB "
        "merupakan batas rawan yang sering memicu sistem crash. Mari kita kalkulasikan penggunaan RAM secara aktual:"
    )
    add_bullet("Windows 11 sistem murni saat idle mengonsumsi sekitar 3.2 GB hingga 3.8 GB RAM.", "• Sistem Operasi Windows: ")
    add_bullet("IDE saat membuka proyek Jetpack Compose atau XML kompleks membutuhkan rata-rata 2.5 GB hingga 3.5 GB RAM.", "• Android Studio IDE: ")
    add_bullet("Ketika melakukan proses Sync dan Build, Gradle Daemon membutuhkan setidaknya 1.5 GB hingga 2.5 GB RAM.", "• Gradle Daemon: ")
    add_bullet("Emulator Pixel yang menjalankan Android 14/15 mengalokasikan RAM virtual mandiri sebesar 2.0 GB hingga 3.0 GB.", "• Android Virtual Device (AVD): ")
    add_bullet("Total kebutuhan riil mencapai 9.2 GB s/d 12.8 GB RAM!", "• AKUMULASI KEBUTUHAN RIIL: ")
    add_p(
        "Jika laptop Anda hanya memiliki RAM 8 GB, sistem operasi Windows terpaksa memindahkan sebagian data memori ke dalam "
        "penyimpanan sekunder melalui mekanisme Virtual Memory (Paging File). Proses pertukaran data inilah yang menyebabkan laptop "
        "mendadak membeku (freeze) dan proses build berlangsung hingga belasan menit."
    )

    add_h3("B. Penyimpanan: Wajib Menggunakan SSD (Solid State Drive)")
    add_p(
        "Penggunaan Harddisk Konvensional (HDD mekanik 5400/7200 RPM) adalah PENYEBAB UTAMA nomor satu lambatnya Android Studio. "
        "Proses kompilasi aplikasi Android melibatkan pembacaan dan penulisan puluhan ribu berkas kecil (I/O intensive). "
        "Pada HDD, disk usage akan langsung melonjak hingga 100%, menyebabkan waktu build memakan waktu 5 hingga 10 menit. "
        "Sebaliknya, pada Solid State Drive (SSD), kecepatan baca-tulis acak ribuan kali lebih tinggi sehingga build yang sama selesai "
        "hanya dalam waktu 15 hingga 30 detik."
    )

    add_h3("C. Processor dan Instruksi Virtualisasi (VT-x / AMD-V)")
    add_p(
        "Processor dengan jumlah Core minimal 4 core dan 8 thread sangat disarankan karena Gradle mengeksekusi kompilasi secara paralel. "
        "Selain itu, fitur Hardware-Assisted Virtualization (Intel VT-x pada prosesor Intel atau AMD-V / SVM pada prosesor AMD) "
        "wajib diaktifkan di dalam pengaturan BIOS/UEFI motherboard. Tanpa fitur ini, emulator Android Virtual Device sama sekali tidak "
        "akan dapat dijalankan."
    )

    add_h2("2.4 Strategi & Tips Optimasi Khusus untuk Laptop RAM 8 GB")
    add_p(
        "Bagi mahasiswa yang memiliki laptop dengan spesifikasi pas-pasan (RAM 8 GB dan processor generasi menengah), "
        "Anda tetap dapat mengikuti praktikum secara lancar dengan menerapkan strategi optimasi berikut:"
    )
    add_bullet(
        "Tinggalkan penggunaan emulator AVD di laptop! Sambungkan smartphone Android fisik Anda menggunakan kabel data USB, "
        "kemudian aktifkan opsi 'USB Debugging' di menu Developer Options HP Anda. Langkah cerdas ini secara instan "
        "menghemat alokasi RAM laptop sebesar 2.5 GB hingga 3.5 GB serta membebaskan beban kerja CPU secara drastis.",
        "1. Gunakan Smartphone Fisik (Physical Device): "
    )
    add_bullet(
        "Buka berkas 'gradle.properties' di dalam direktori proyek Anda atau di 'C:\\Users\\<NamaUser>\\.gradle\\gradle.properties', "
        "kemudian tambahkan baris konfigurasi berikut untuk membatasi konsumsi memori Gradle:\n"
        "org.gradle.jvmargs=-Xmx1536m -XX:MaxMetaspaceSize=512m\n"
        "Pengaturan ini membatasi alokasi maksimal heap Gradle agar tidak melebihi 1.5 GB.",
        "2. Batasi Heap Memory Gradle JVM: "
    )
    add_bullet(
        "Saat sedang mengetik baris kode dan merancang antarmuka, aktifkan mode 'Power Save Mode' melalui menu "
        "'File > Power Save Mode'. Fitur ini akan menonaktifkan inspeksi syntax background yang terus menerus memakan siklus CPU.",
        "3. Manfaatkan Fitur Power Save Mode: "
    )
    add_bullet(
        "Selalu tutup aplikasi berat lain seperti Google Chrome (terutama tab video YouTube), Discord, Spotify, dan "
        "game launcher sebelum membuka Android Studio untuk membebaskan ruang memori.",
        "4. Matikan Aplikasi Latar Belakang: "
    )

    add_callout(
        "Bila laptop Anda masih menggunakan konfigurasi RAM 8 GB Single Channel, sangat disarankan untuk melakukan "
        "upgrade mandiri ke 16 GB Dual Channel (menambah 1 keping RAM 8 GB serupa). Biaya upgrade RAM relatif sangat terjangkau "
        "dibandingkan peningkatan produktivitas yang akan Anda rasakan selama menempuh semester 7 dan pengerjaan Tugas Akhir.",
        title="REKOMENDASI UPGRADE HARDWARE",
        box_type="warning"
    )

    # =========================================================================
    # BAB III: PANDUAN LANGKAH-LANGKAH PENGUNDUHAN DAN INSTALASI
    # =========================================================================
    add_h1("BAB III: LANGKAH-LANGKAH PENGUNDUHAN DAN INSTALASI")
    add_p(
        "Proses penginstalan Android Studio terbagi menjadi empat fase utama yang berurutan:\n"
        "1. Fase 1: Pengunduhan Berkas Installer Resmi (Langkah 1)\n"
        "2. Fase 2: Pemasangan Perangkat Lunak Inti pada Sistem Windows (Langkah 2 s/d Langkah 6)\n"
        "3. Fase 3: Inisialisasi Pertama & Setup Android SDK Toolchain (Langkah 7 s/d Langkah 14)\n"
        "4. Fase 4: Dashboard Utama Selamat Datang Siap Digunakan (Langkah 15)\n\n"
        "Berikut adalah penjelasan rinci untuk setiap langkah yang disertai dengan screenshot otentik dari proses instalasi."
    )

    # -------------------------------------------------------------------------
    # FASE 1
    # -------------------------------------------------------------------------
    add_h2("FASE 1: PENGUNDUHAN BERKAS INSTALLER RESMI")
    
    add_h3("Langkah 1: Mengunduh Installer Resmi dari Portal Android Developers")
    add_figure(
        "Step_01_Download_Android_Studio.png",
        "Gambar 1: Halaman Resmi Pengunduhan Android Studio di developer.android.com/studio",
        width=Inches(5.6)
    )
    add_p(
        "Buka peramban (web browser) di laptop Anda dan akses situs web resmi pengembang Android pada tautan: "
        "https://developer.android.com/studio. Situs ini secara otomatis mendeteksi sistem operasi yang Anda gunakan (Windows 64-bit) "
        "dan menampilkan tombol utama pengunduhan bertuliskan 'Download Android Studio Quail 4'.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Klik tombol unduh tersebut. Anda akan diminta membaca dan menyetujui lembar syarat dan ketentuan (Terms and Conditions). "
        "Beri centang pada kotak persetujuan 'I have read and agree with the above terms and conditions', kemudian klik tombol unduh warna biru. "
        "Berkas installer yang diunduh bernama 'android-studio-quail4-windows.exe' dengan ukuran sekitar 1.4 GB. Simpan berkas tersebut "
        "ke dalam folder Downloads laptop Anda dan tunggu hingga proses unduhan selesai 100%.",
        bold_prefix="Rincian Teknis: "
    )

    # -------------------------------------------------------------------------
    # FASE 2
    # -------------------------------------------------------------------------
    add_h2("FASE 2: PEMASANGAN PERANGKAT LUNAK INTI (WINDOWS SETUP WIZARD)")

    add_h3("Langkah 2: Menjalankan Installer & Layar Sambutan (Welcome Setup)")
    add_figure(
        "Step_02_Installer_Welcome.png",
        "Gambar 2: Layar Awal Selamat Datang pada Android Studio Setup Wizard",
        width=Inches(4.8)
    )
    add_p(
        "Buka folder Downloads, cari berkas 'android-studio-quail4-windows.exe', klik kanan pada berkas tersebut lalu pilih "
        "'Run as administrator'. Apabila muncul jendela peringatan User Account Control (UAC) dari Windows, pilih 'Yes'. "
        "Jendela panduan awal 'Welcome to Android Studio Setup' akan terbuka di layar Anda.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Layar ini merekomendasikan pengguna untuk menutup aplikasi lain yang sedang berjalan agar instalasi dapat memperbarui "
        "berkas sistem terkait tanpa harus melakukan restart komputer. Klik tombol 'Next >' untuk melangkah ke tahap pemilihan komponen.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 3: Pemilihan Komponen Perangkat Lunak (Choose Components)")
    add_figure(
        "Step_03_Choose_Components.png",
        "Gambar 3: Pemilihan Komponen Instalasi Android Studio & Android Virtual Device",
        width=Inches(4.8)
    )
    add_p(
        "Pada halaman 'Choose Components', Anda disajikan daftar fitur yang akan dipasang ke dalam sistem. Terdapat dua pilihan komponen utama: "
        "'Android Studio' (inti IDE perangkat lunak) dan 'Android Virtual Device' (perangkat virtual emulator bawaan).",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Komponen 'Android Studio' berstatus wajib dan kotak centangnya terkunci abu-abu. Pastikan kotak centang kedua yaitu "
        "'Android Virtual Device' (AVD) juga TERCENTANG. Komponen AVD ini sangat penting karena menyediakan konfigurasi emulator "
        "untuk menjalankan ponsel cerdas virtual di PC Anda. Total ruang disk yang dibutuhkan pada tahap ini sekitar 3.8 GB. "
        "Setelah memastikan kedua opsi tercentang, klik tombol 'Next >'.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 4: Konfigurasi Lokasi Direktori Instalasi (Install Location)")
    add_figure(
        "Step_04_Install_Location.png",
        "Gambar 4: Menentukan Direktori Tujuan Pemasangan Android Studio",
        width=Inches(4.8)
    )
    add_p(
        "Pada jendela 'Configuration Settings: Install Locations', Anda diminta menentukan jalur folder direktori tempat seluruh program "
        "biner Android Studio akan diletakkan di harddisk / SSD laptop.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Jalur default yang disediakan oleh sistem adalah 'C:\\Program Files\\Android\\Android Studio'. Sangat disarankan untuk membiarkan "
        "jalur lokasi ini tetap pada kondisi default-nya demi menghindari masalah perizinan (permissions) atau path environment variables. "
        "Pastikan partisi drive C: Anda masih memiliki ruang kosong yang memadai (minimal 10 GB ruang bebas). Klik tombol 'Next >', "
        "lalu pada jendela pilihan folder Start Menu langsung klik tombol 'Install' untuk memulai proses pemasangan.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 5: Proses Ekstraksi dan Pemasangan Berkas (Installing)")
    add_figure(
        "Screenshot 2026-09-07 100427.png",
        "Gambar 5: Proses Ekstraksi Berkas Sistem dan Pemasangan Komponen Berlangsung",
        width=Inches(4.8)
    )
    add_p(
        "Sistem Windows mulai mengekstrak seluruh arsip paket program Android Studio ke direktori tujuan. Layar menampilkan bilah progres "
        "hijau beserta nama berkas yang sedang diekstrak secara real-time (contoh: 'Extract: NotoSansCJK-Regular.ttc...').",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Pada tahap ini, berkas yang dipasang mencakup engine IntelliJ IDEA, font antarmuka multi-bahasa, utilitas decompilation, "
        "serta runtime Java khusus JetBrains Runtime (JBR) yang dioptimalkan untuk performa IDE. Anda dapat menekan tombol 'Show details' "
        "apabila ingin melihat log berkas yang diekstrak satu demi satu. Tunggu hingga bilah progres berwarna hijau penuh 100% dan tombol "
        "'Next >' di bagian bawah menjadi aktif (dapat diklik).",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 6: Penyelesaian Instalasi Perangkat Lunak Inti (Completing Setup)")
    add_figure(
        "Screenshot 2026-09-07 100532.png",
        "Gambar 6: Konfirmasi Penyelesaian Instalasi Utama Android Studio",
        width=Inches(4.8)
    )
    add_p(
        "Jendela 'Completing Android Studio Setup' mengonfirmasikan bahwa perangkat lunak Android Studio telah berhasil terpasang sempurna "
        "pada komputer Anda ('Android Studio has been installed on your computer').",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Pada halaman ini terdapat kotak centang bertuliskan 'Start Android Studio'. Pastikan kotak tersebut DIBIARKAN TERCENTANG agar "
        "aplikasi langsung dijalankan secara otomatis saat Anda menutup jendela installer. Selanjutnya, klik tombol 'Finish' di bagian kanan bawah.",
        bold_prefix="Rincian Teknis: "
    )

    # -------------------------------------------------------------------------
    # FASE 3
    # -------------------------------------------------------------------------
    add_h2("FASE 3: INISIALISASI PERTAMA & SETUP ANDROID SDK TOOLCHAIN")

    add_h3("Langkah 7: Menjalankan Android Studio & Splash Screen")
    add_figure(
        "Screenshot 2026-09-07 100551.png",
        "Gambar 7: Splash Screen Inisialisasi Android Studio Quail 4 (Rilis 2026.1.4)",
        width=Inches(5.2)
    )
    add_p(
        "Setelah menekan tombol Finish, sistem akan meluncurkan Android Studio. Layar pembuka (splash screen) khas Google akan muncul "
        "menampilkan ilustrasi burung puyuh mahkota (Quail) yang menjadi kode nama resmi versi Quail 4 (2026.1.4), didukung platform IntelliJ.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Di balik layar, Android Studio sedang melakukan inisialisasi modul Java Virtual Machine (JVM), memuat pustaka plug-in inti, "
        "dan memeriksa apakah komputer sudah memiliki konfigurasi Android SDK sebelumnya atau merupakan instalasi baru murni (clean install).",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 8: Kebijakan Privasi & Telemetri (Help Improve Android Studio)")
    add_figure(
        "Screenshot 2026-09-07 100619.png",
        "Gambar 8: Dialog Pengaturan Berbagi Data Penggunaan dan Statistik Telemetri",
        width=Inches(4.6)
    )
    add_p(
        "Sebuah jendela dialog bertajuk 'Help improve Android Studio' akan muncul di tengah layar. Google meminta izin kepada pengguna "
        "apakah Anda bersedia mengirimkan statistik penggunaan fitur, konsumsi resource, dan laporan error (crash log) secara anonim.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Tersedia dua opsi pilihan:\n"
        "• 'Send usage statistics': Menyetujui pengiriman data diagnostik telemetri anonim ke server Google.\n"
        "• 'Don't send': Menolak pengiriman data telemetri penggunaan.\n"
        "Kedua pilihan ini SAMA SEKALI TIDAK MEMPENGARUHI fitur maupun kinerja IDE. Anda bebas memilih 'Don't send' untuk menghemat "
        "privasi dan lalu lintas bandwidth latar belakang, atau 'Send usage statistics' jika ingin berkontribusi pada pengembangan ekosistem Android.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 9: Memulai Android Studio Setup Wizard (First-Run Environment)")
    add_figure(
        "Screenshot 2026-09-07 100640.png",
        "Gambar 9: Layar Selamat Datang Wizard Pengaturan Lingkungan Pengembangan (SDK Setup)",
        width=Inches(5.0)
    )
    add_p(
        "Karena ini adalah instalasi pertama kali pada komputer Anda, wizard khusus 'Android Studio Setup Wizard' akan secara otomatis "
        "terbuka menyapa Anda: 'Welcome! This wizard will set up your development environment for Android Studio...'.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Wizard ini memegang peran paling krusial dalam mengubah Android Studio dari sekadar text editor menjadi lingkungan pengembangan mobile "
        "yang lengkap, karena wizard inilah yang akan mengunduh toolchain Android SDK, kompiler platform, dan emulator AVD. "
        "Klik tombol 'Next' pada bagian kanan bawah untuk melangkah ke konfigurasi setup.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 10: Pemilihan Tipe Instalasi (Install Type: Standard vs Custom)")
    add_figure(
        "Screenshot 2026-09-07 100650.png",
        "Gambar 10: Memilih Tipe Setup Android Studio (Standard vs Custom)",
        width=Inches(5.0)
    )
    add_p(
        "Pada halaman 'Install Type', Anda dihadapkan pada dua mode konfigurasi:\n"
        "1. Standard: Android Studio akan diinstal dengan paket pengaturan dan komponen yang paling umum digunakan (sangat disarankan bagi sebagian besar pengguna).\n"
        "2. Custom: Memungkinkan Anda mengatur sendiri lokasi folder SDK, alokasi RAM emulator, dan memilih versi komponen individual.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Untuk keperluan praktikum perkuliahan PAM, PILIHLAH OPSI 'Standard' (pastikan radio button Standard tercentang biru). "
        "Pilihan Standard menjamin bahwa versi SDK Platform, Build-Tools, dan emulator yang diunduh adalah paket yang paling stabil "
        "dan kompatibel satu sama lain. Setelah memilih Standard, klik tombol 'Next'.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 11: Verifikasi Pengaturan Komponen SDK (Verify Settings)")
    add_figure(
        "Screenshot 2026-09-07 100700.png",
        "Gambar 11: Jendela Verifikasi Daftar Komponen SDK yang Akan Diunduh",
        width=Inches(5.0)
    )
    add_p(
        "Layar 'Verify Settings' menampilkan daftar rekapitulasi komponen Android SDK yang akan diunduh dan dipasang oleh wizard ke laptop Anda. "
        "Tinjau informasi yang tertera sebelum melanjutkan.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Informasi penting yang ditampilkan meliputi:\n"
        "• Setup Type: Standard\n"
        "• SDK Folder: C:\\Users\\<NamaUser>\\AppData\\Local\\Android\\Sdk (lokasi instalasi SDK bawaan Windows)\n"
        "• Total Download Size: Sekitar 600 MB (arsip terkompresi)\n"
        "• Rincian Komponen SDK yang akan dipasang:\n"
        "  1. Android Emulator: 421 MB (mesin virtual perangkat)\n"
        "  2. Android SDK Build-Tools 36: 56 MB (kompiler dan packaging tool)\n"
        "  3. Android SDK Platform 37.0: 64.2 MB (library inti Android framework)\n"
        "  4. Android SDK Platform-Tools: 7.67 MB (utilitas ADB dan Fastboot)\n"
        "  5. Sources for Android 37.0: 50.4 MB (kode sumber dokumentasi class Android)\n"
        "Pastikan koneksi internet laptop Anda aktif, kemudian klik tombol 'Next'.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 12: Persetujuan Perjanjian Lisensi (License Agreement)")
    add_figure(
        "Screenshot 2026-09-07 100713.png",
        "Gambar 12: Persetujuan Lisensi Legal Android Software Development Kit (SDK)",
        width=Inches(5.0)
    )
    add_p(
        "Pada halaman 'License Agreement', Anda diwajibkan menyetujui perjanjian lisensi penggunaan perangkat lunak resmi dari Google "
        "('Android Software Development Kit License Agreement') sebelum proses pengunduhan diizinkan dimulai.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Langkah yang harus dilakukan:\n"
        "1. Klik pada nama kelompok lisensi 'android-sdk-license' di panel sebelah kiri.\n"
        "2. Perhatikan bagian bawah panel kanan, klik radio button bertuliskan 'Accept' (berubah menjadi bulatan biru).\n"
        "3. (Catatan: Apabila terdapat lisensi tambahan seperti 'android-arm-dbt-license' atau 'intel-android-extra-license' pada daftar kiri, klik masing-masing lisensi tersebut dan pilih 'Accept').\n"
        "4. Setelah seluruh lisensi berstatus disetujui, tombol 'Finish' di pojok kanan bawah akan menjadi aktif dan menyala biru. Klik tombol 'Finish'.",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 13: Proses Pengunduhan Komponen SDK (Downloading Components)")
    add_figure(
        "Screenshot 2026-09-07 100722.png",
        "Gambar 13: Pengunduhan Berkas Arsip SDK dari Repositori Resmi Google",
        width=Inches(5.0)
    )
    add_p(
        "Wizard mulai menghubungi server repositori resmi Google (melalui URL https://dl.google.com/android/repository/...) "
        "untuk mengunduh satu demi satu paket komponen yang telah diverifikasi pada langkah sebelumnya.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Layar menampilkan nama berkas arsip yang sedang diunduh (contoh: 'Downloading source-37.0_r02.zip (11%): 5.3 / 48.2 MB ...') "
        "beserta persentase progres. Kecepatan pada tahap ini sepenuhnya bergantung pada kestabilan dan bandwidth internet Anda. "
        "JANGAN MENUTUP JENDELA INI atau memutuskan koneksi internet selama proses unduh berlangsung demi mencegah terjadinya berkas arsip korup (corrupt zip).",
        bold_prefix="Rincian Teknis: "
    )

    add_h3("Langkah 14: Pengunduhan SDK Selesai dan Terverifikasi")
    add_figure(
        "Screenshot 2026-09-07 101547.png",
        "Gambar 14: Log Konfirmasi Bahwa Seluruh Komponen SDK Telah Terpasang Sempurna",
        width=Inches(5.0)
    )
    add_p(
        "Setelah seluruh unduhan selesai, wizard akan secara otomatis mengekstrak berkas zip ke dalam folder SDK lokal Anda. "
        "Jendela log akan menampilkan ringkasan instalasi secara mendetail.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Perhatikan baris-baris log penting berikut:\n"
        "• 'Installing Android SDK Platform 37.0 ... complete & finished'\n"
        "• 'Installing Android SDK Platform-Tools v.37.0.1 ... complete & finished'\n"
        "• 'Installing Android Emulator v.37.1.11 ... complete & finished'\n"
        "• 'SDK Manager found the following installed packages: build-tools;36.0.0, emulator, platform-tools, platforms;android-37.0, sources;android-37.0'\n"
        "• 'Android SDK is up to date.'\n"
        "Pesan 'Android SDK is up to date' menandakan bahwa seluruh toolchain inti pengembangan Android telah berhasil dipasang tanpa ada cacat berkas. "
        "Klik tombol 'Finish' di bagian kanan bawah untuk mengakhiri setup wizard.",
        bold_prefix="Rincian Teknis: "
    )

    # -------------------------------------------------------------------------
    # FASE 4
    # -------------------------------------------------------------------------
    add_h2("FASE 4: DASHBOARD UTAMA SIAP DIGUNAKAN")

    add_h3("Langkah 15: Tampilan Selamat Datang Android Studio (Welcome Screen)")
    add_figure(
        "Screenshot 2026-09-07 101602.png",
        "Gambar 15: Dashboard Utama Selamat Datang di Android Studio Quail 4",
        width=Inches(5.0)
    )
    add_p(
        "Selamat! Lingkungan pengembangan Android Studio kini telah siap 100% digunakan. Anda akan disambut oleh dashboard utama "
        "bertajuk 'Welcome to Android Studio'.",
        bold_prefix="Deskripsi Tindakan: "
    )
    add_p(
        "Pada dashboard ini terdapat beberapa opsi penting yang akan kita gunakan sepanjang perkuliahan PAM:\n"
        "1. Panel Aksi Utama (Tengah):\n"
        "   • 'New Project': Membuat proyek aplikasi Android baru dari awal menggunakan berbagai template yang disediakan.\n"
        "   • 'Open': Membuka proyek Android yang sudah tersimpan di direktori penyimpanan komputer Anda.\n"
        "   • 'Clone Repository': Mengklon repositori proyek secara langsung dari platform Git remote (seperti GitHub atau GitLab).\n"
        "2. Panel Navigasi Samping (Kiri):\n"
        "   • 'Projects': Menampilkan daftar riwayat proyek yang baru saja dibuka (recent projects).\n"
        "   • 'Customize': Mengatur preferensi antarmuka visual IDE (tema Gelap/Terang, ukuran huruf editor font, dan keymaps).\n"
        "   • 'Plugins': Mengelola ekstensi tambahan untuk memperkaya kapabilitas IDE.\n"
        "   • 'Learn': Akses ke tutorial resmi Google untuk mempelajari sintaks dasar dan fitur baru Android.",
        bold_prefix="Rincian Teknis: "
    )

    add_callout(
        "Sampai di titik ini, komputer Anda telah resmi memiliki lingkungan pengembangan aplikasi mobile berstandar industri. "
        "Pastikan Anda menyimpan jalur folder SDK ('C:\\Users\\<NamaUser>\\AppData\\Local\\Android\\Sdk') dan tidak menghapus folder "
        "tersebut secara tidak sengaja.",
        title="INSTALASI SELESAI DENGAN SUKSES",
        box_type="tip"
    )

    # =========================================================================
    # BAB IV: VERIFIKASI INSTALASI & PENGUJIAN PROJECT PERTAMA
    # =========================================================================
    add_h1("BAB IV: VERIFIKASI INSTALASI & PENGUJIAN PROJECT PERTAMA")
    add_p(
        "Untuk memastikan bahwa Android Studio, compiler Kotlin, Android SDK, dan sistem build Gradle bekerja secara harmonis, "
        "setiap mahasiswa diwajibkan melakukan uji coba pembuatan proyek perdana (Hello World) dengan mengikuti langkah-langkah berikut:"
    )

    add_h2("4.1 Langkah Pembuatan Proyek Baru (New Project)")
    add_bullet("Pada dashboard utama (Gambar 15), klik tombol '+ New Project'.", "1. Buka Wizard Proyek: ")
    add_bullet(
        "Pilih kategori 'Phone and Tablet' di panel kiri, kemudian pilih template 'Empty Views Activity' "
        "(jika mata kuliah menggunakan pendekatan XML Klasik) atau 'Empty Activity' (jika mata kuliah menggunakan Jetpack Compose modern). "
        "Klik tombol 'Next'.",
        "2. Pemilihan Template: "
    )
    add_bullet(
        "Isi parameter proyek sebagai berikut:\n"
        "   • Name: 'MyFirstPAMApp'\n"
        "   • Package name: 'com.example.myfirstpamapp' (penanda unik aplikasi di ekosistem Android)\n"
        "   • Save location: Tentukan folder di partisi penyimpanan Anda, misal 'c:\\Kuliah\\Semester 7\\PAM\\MyFirstPAMApp'\n"
        "   • Language: Pilih 'Kotlin' (bahasa resmi utama yang diwajibkan oleh Google)\n"
        "   • Minimum SDK: Pilih 'API 26: Android 8.0 (Oreo)' atau 'API 24: Android 7.0 (Nougat)'. "
        "Pilihan ini memungkinkan aplikasi Anda dapat berjalan pada lebih dari 95% perangkat Android aktif di seluruh dunia.\n"
        "   • Build configuration language: 'Kotlin DSL (build.gradle.kts)'.",
        "3. Konfigurasi Proyek: "
    )
    add_bullet("Klik tombol 'Finish' di kanan bawah. Android Studio akan mulai merangkai struktur berkas proyek perdana Anda.", "4. Inisialisasi: ")

    add_h2("4.2 Menunggu Proses Gradle Sync Selesai")
    add_p(
        "Ketika editor terbuka untuk pertama kalinya, perhatikan bilah status (status bar) di bagian paling bawah jendela IDE. "
        "Android Studio sedang menjalankan proses 'Gradle Sync' dan mengunduh dependensi pustaka dasar dari Google Maven Repository. "
        "Proses ini membutuhkan koneksi internet aktif selama 1 hingga 3 menit. "
        "Tanda keberhasilan adalah munculnya pesan: 'Gradle sync finished in ... s' dan struktur pohon direktori di tab 'Project' "
        "menampilkan folder 'app', 'manifests', 'java/kotlin', dan 'res'."
    )

    add_h2("4.3 Menjalankan Aplikasi (Run App)")
    add_p("Untuk mengeksekusi aplikasi ke perangkat pengujian, Anda dapat memilih salah satu dari dua metode:")
    add_bullet(
        "Buka Device Manager di bilah kanan atas Android Studio. Klik tanda play (Run) pada virtual device yang telah terinstal. "
        "Setelah emulator ponsel muncul di layar, pilih nama emulator tersebut di dropdown target device di toolbar atas, "
        "kemudian klik tombol 'Run 'app'' (ikon Segitiga Hijau atau tekan tombol Shift + F10).",
        "A. Menggunakan Android Virtual Device (Emulator): "
    )
    add_bullet(
        "Hubungkan ponsel Android Anda ke laptop dengan kabel data USB. Pastikan opsi 'USB Debugging' pada menu Developer Options di HP aktif. "
        "Ketika muncul popup 'Allow USB debugging?' di layar HP, centang 'Always allow from this computer' lalu ketuk 'OK'. "
        "Nama ponsel Anda akan langsung terdeteksi di dropdown device Android Studio. Klik tombol Segitiga Hijau 'Run 'app''. "
        "Aplikasi akan otomatis terkompilasi, dikirim, dan terbuka di layar smartphone Anda menampilkan teks 'Hello Android!'.",
        "B. Menggunakan HP Android Fisik (Sangat Disarankan untuk RAM 8 GB): "
    )

    # =========================================================================
    # BAB V: PANDUAN PEMECAHAN MASALAH (TROUBLESHOOTING GUIDE)
    # =========================================================================
    add_h1("BAB V: PANDUAN PEMECAHAN MASALAH (TROUBLESHOOTING GUIDE)")
    add_p(
        "Berikut adalah rangkuman solusi atas lima kendala teknis paling umum yang sering dihadapi mahasiswa pada pertemuan perdana:"
    )

    add_h2("5.1 Error: 'VT-x / AMD-V is not enabled in BIOS / Hypervisor Error'")
    add_p(
        "Gejala: Saat mencoba meluncurkan emulator AVD, muncul pesan peringatan: 'Android Emulator requires hardware acceleration "
        "such as VT-x / SVM which is disabled in BIOS'."
    )
    add_bullet(
        "Restart laptop Anda, dan segera tekan tombol tombol BIOS (biasanya F2, F10, F12, atau Del tergantung merek laptop seperti ASUS, Lenovo, HP, Acer, Dell) "
        "secara berulang-ulang sebelum logo Windows muncul.\n"
        "Masuk ke menu 'Advanced' atau 'CPU Configuration', cari pengaturan bertuliskan 'Intel Virtualization Technology' (untuk Intel) "
        "atau 'SVM Mode / AMD-V' (untuk AMD). Ubah nilainya menjadi 'Enabled'. Tekan F10 untuk menyimpan dan restart ke Windows.",
        "Solusi 1 (BIOS): "
    )
    add_bullet(
        "Buka Start Menu Windows, ketik 'Turn Windows features on or off', buka menu tersebut. Beri tanda centang pada opsi "
        "'Windows Hypervisor Platform' dan 'Virtual Machine Platform'. Klik OK lalu restart laptop Anda.",
        "Solusi 2 (Fitur Windows): "
    )

    add_h2("5.2 Error: 'Gradle Sync Failed: Connection Timed Out / PKIX Path Building Failed'")
    add_p(
        "Gejala: Gradle sync gagal dengan tulisan merah tebal 'Could not resolve com.android.tools.build:gradle...' atau 'Connection timed out'."
    )
    add_bullet(
        "Masalah ini 95% disebabkan oleh koneksi internet yang tidak stabil, terhalang firewall kampus / kantor, atau adanya proxy. "
        "Solusinya: Beralihlah sementara ke koneksi Hotspot seluler (Tethering HP pribadi). "
        "Setelah tersambung ke hotspot yang lancar, klik tombol 'Sync Project with Gradle Files' (ikon Gajah dengan panah biru di toolbar kanan atas).",
        "Solusi: "
    )

    add_h2("5.3 Error: 'Your SDK Path Contains Whitespace / Non-ASCII Characters'")
    add_p(
        "Gejala: Timbul peringatan bahwa folder SDK mengandung karakter spasi atau karakter khusus (misal nama akun Windows Anda: "
        "'C:\\Users\\Budi Santoso\\AppData\\Local\\Android\\Sdk'). Beberapa alat biner SDK (seperti NDK dan AAPT2) gagal memproses path dengan spasi."
    )
    add_bullet(
        "Buat folder baru di drive C: tanpa spasi, contoh: 'C:\\Android\\Sdk'. Buka Android Studio, masuk ke menu "
        "'Tools > SDK Manager'. Pada bagian 'Android SDK Location', klik tombol 'Edit' dan arahkan ke folder baru 'C:\\Android\\Sdk'. "
        "Biarkan Android Studio mengunduh atau menyalin komponen ke folder baru tersebut.",
        "Solusi: "
    )

    add_h2("5.4 Error: 'Device Unauthorized / Smartphone Fisik Tidak Terdeteksi'")
    add_p(
        "Gejala: Smartphone sudah dicolok menggunakan kabel data USB tetapi tidak muncul di dropdown device Android Studio, "
        "atau muncul dengan keterangan 'Unauthorized'."
    )
    add_bullet(
        "Pastikan kabel yang digunakan adalah kabel data (bukan kabel charger murah yang hanya menghantarkan daya listrik tanpa jalur data). "
        "Ubah mode koneksi USB di HP dari 'Charging only' menjadi 'File Transfer (MTP)'. "
        "Di HP, buka Settings > Developer Options > matikan lalu hidupkan kembali opsi 'USB Debugging', lalu ketuk 'Revoke USB debugging authorizations'. "
        "Cabut dan tancapkan kembali kabel USB, lalu pastikan Anda mengetuk 'Allow' pada dialog popup otorisasi di layar HP.",
        "Solusi: "
    )

    add_h2("5.5 Error: 'Gradle Out of Memory / Java Heap Space'")
    add_p(
        "Gejala: Proses kompilasi terhenti dengan pesan error: 'java.lang.OutOfMemoryError: Java heap space'."
    )
    add_bullet(
        "Buka berkas 'gradle.properties' yang terletak di root direktori proyek Anda. Temukan atau tambahkan baris berikut:\n"
        "org.gradle.jvmargs=-Xmx2048m -XX:MaxMetaspaceSize=512m\n"
        "Jika laptop Anda memiliki RAM 16 GB, Anda dapat menaikkannya menjadi '-Xmx4096m' untuk mempercepat proses kompilasi.",
        "Solusi: "
    )

    # =========================================================================
    # BAB VI: KESIMPULAN & CHECKLIST KESIAPAN
    # =========================================================================
    add_h1("BAB VI: KESIMPULAN & CHECKLIST KESIAPAN")
    add_p(
        "Penyelesaian seluruh tahapan instalasi Android Studio Quail 4 (2026.1.4) merupakan pencapaian penting dalam memulai mata kuliah "
        "Pengembangan Aplikasi Mobile (PAM). Dengan lingkungan pengembangan yang terkonfigurasi secara tepat, mahasiswa kini memiliki toolchain "
        "lengkap untuk merancang, mengodekan, menguji, dan mendebug aplikasi Android modern."
    )
    add_p(
        "Sebelum menghadiri Pertemuan ke-2 pekan depan, setiap mahasiswa diwajibkan melakukan verifikasi mandiri menggunakan lembar checklist berikut:"
    )

    # Checklist Table
    tbl_chk = doc.add_table(rows=11, cols=3)
    tbl_chk.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_chk, color="BDC3C7", sz="4")
    
    chk_headers = ["No", "Item Verifikasi Kelayakan", "Status Kesiapan"]
    for i, h in enumerate(chk_headers):
        cell = tbl_chk.rows[0].cells[i]
        set_cell_background(cell, "073042")
        set_cell_margins(cell, top=140, bottom=140, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.bold = True
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    checklist_items = [
        ("1", "Berkas installer Android Studio Quail 4 (2026.1.4) berhasil dipasang tanpa error.", "[   ] Siap"),
        ("2", "Android Studio dapat diluncurkan dan menampilkan dashboard 'Welcome to Android Studio'.", "[   ] Siap"),
        ("3", "Android SDK Platform 37.0 terpasang di C:\\Users\\<user>\\AppData\\Local\\Android\\Sdk.", "[   ] Siap"),
        ("4", "Android SDK Build-Tools versi 36.0.0 dan Platform-Tools berhasil terinstal.", "[   ] Siap"),
        ("5", "Lisensi 'android-sdk-license' telah disetujui (Accept) dan berstatus valid.", "[   ] Siap"),
        ("6", "Fitur Hardware Virtualization (VT-x / AMD-V) telah diaktifkan di BIOS laptop.", "[   ] Siap"),
        ("7", "Tersedia ruang kosong (free space) minimal 15-20 GB pada drive SSD sistem.", "[   ] Siap"),
        ("8", "Proyek uji coba 'MyFirstPAMApp' berhasil dibuat dan proses Gradle Sync sukses tanpa error.", "[   ] Siap"),
        ("9", "Perangkat pengujian (Android Virtual Device atau Smartphone Fisik via USB) berfungsi normal.", "[   ] Siap"),
        ("10", "Aplikasi 'Hello World' berhasil dijalankan dan tampil sempurna di layar perangkat pengujian.", "[   ] Siap")
    ]

    for row_idx, (no, item, status) in enumerate(checklist_items, start=1):
        bg = "F8F9FA" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate([no, item, status]):
            cell = tbl_chk.rows[row_idx].cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            if col_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.bold = True
            elif col_idx == 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run.bold = True
                run.font.color.rgb = RGBColor(46, 125, 50)

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_before = Pt(20)
    p_end.paragraph_format.space_after = Pt(0)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_end = p_end.add_run("— SEMOGA PRAKTIKUM ANDA BERJALAN LANCAR DAN SUKSES —")
    run_end.bold = True
    run_end.font.name = "Arial"
    run_end.font.size = Pt(10)
    run_end.font.color.rgb = RGBColor(120, 125, 130)

    # Save documents with both official names
    filename1 = "Panduan_Instalasi_Android_Studio_dan_Spesifikasi_Laptop.docx"
    filename2 = "Muhammad Fatahillah Farid_123140203_P1.docx"
    doc.save(filename1)
    doc.save(filename2)
    print(f"Documents successfully created:\n  - {filename1}\n  - {filename2}")

if __name__ == "__main__":
    create_document()
